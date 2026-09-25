import argparse
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = '1F4E79'; DARK = '0B2545'; LIGHT_BLUE = 'D9EAF7'; LIGHT_GRAY = 'F2F4F7'
PALE_YELLOW = 'FFF2CC'; PALE_RED = 'FCE4D6'; PALE_GREEN = 'E2F0D9'; BORDER = 'D9E2F3'


def style_run(run, bold=False, size=None, color=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color='D9E2F3', size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0'); element.set(qn('w:color'), color)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)
    sec.header_distance = Inches(0.45); sec.footer_distance = Inches(0.45)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in [('Heading 1', 15, BLUE, 12, 6), ('Heading 2', 12.5, BLUE, 8, 4), ('Heading 3', 11.5, DARK, 6, 3)]:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    for name in ['List Bullet', 'List Number']:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.12


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        style_run(r, bold=True, color=BLUE if level < 3 else DARK)


def callout(doc, title, lines, fill=LIGHT_BLUE):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; set_table_borders(t, color=BORDER)
    c = t.cell(0, 0); set_cell_shading(c, fill); set_cell_margins(c, top=130, bottom=130, start=160, end=160)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); style_run(r, bold=True, size=12, color=DARK)
    for line in lines or []:
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(line)); style_run(r, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, first_col_fill=None, font_size=9.5):
    if not rows:
        rows = [['待补充' for _ in headers]]
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; set_table_borders(t); set_table_widths(t, widths)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; set_cell_shading(c, header_fill)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); style_run(r, bold=True, size=font_size, color=DARK)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if first_col_fill and i == 0:
                set_cell_shading(cells[i], first_col_fill)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val)); style_run(r, size=font_size, bold=(first_col_fill and i == 0), color=DARK if i == 0 else None)
    set_table_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def numbers(doc, items):
    for item in items or ['待补充']:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(str(item)); style_run(r)


def build(data, output):
    doc = Document(); configure(doc)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(data.get('title', '需求分析报告')); style_run(r, bold=True, size=22, color=DARK)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(data.get('metadata', '')); style_run(r, size=10, color='666666')

    callout(doc, '一页结论', data.get('conclusion', []), fill=LIGHT_BLUE)
    heading(doc, '规则速览', 1)
    add_table(doc, ['clientId/对象', '判断指标/规则', '统计周期', '阈值/条件', '满足条件'],
              [[x.get('clientId', ''), x.get('metric', ''), x.get('period', ''), x.get('threshold', ''), x.get('condition', '')] for x in data.get('rule_overview', [])],
              [Inches(1.05), Inches(2.2), Inches(1.1), Inches(0.9), Inches(1.55)], header_fill=LIGHT_BLUE, font_size=10)
    heading(doc, '优先确认项', 1)
    add_table(doc, ['优先级', '待确认问题', '不确认的影响', '确认结论'],
              [[x.get('priority', ''), x.get('question', ''), x.get('impact', ''), x.get('confirmation', '待填写')] for x in data.get('priority_items', [])],
              [Inches(0.55), Inches(2.9), Inches(2.0), Inches(1.65)], header_fill=PALE_YELLOW, first_col_fill=PALE_YELLOW, font_size=8.6)

    doc.add_page_break()
    heading(doc, '详细疑问清单', 1); numbers(doc, data.get('detailed_questions', []))
    heading(doc, '模糊点与澄清建议', 1)
    add_table(doc, ['模糊点', '问题说明', '建议澄清'],
              [[x.get('term', ''), x.get('issue', ''), x.get('suggestion', '')] for x in data.get('ambiguities', [])],
              [Inches(1.35), Inches(2.25), Inches(3.0)], font_size=9.2)
    heading(doc, '验收标准建议', 1)
    add_table(doc, ['编号', '验收场景', '预期结果'],
              [[x.get('id', ''), x.get('scenario', ''), x.get('expected', '')] for x in data.get('acceptance_criteria', [])],
              [Inches(0.7), Inches(4.15), Inches(1.75)], header_fill=PALE_GREEN, first_col_fill=PALE_GREEN, font_size=9.2)
    heading(doc, '风险提示', 1)
    add_table(doc, ['风险等级', '风险点', '说明'],
              [[x.get('level', ''), x.get('risk', ''), x.get('description', '')] for x in data.get('risks', [])],
              [Inches(0.8), Inches(2.0), Inches(3.8)], header_fill=PALE_RED, first_col_fill=PALE_RED, font_size=9.3)
    heading(doc, '结论', 1); callout(doc, '开发前建议', data.get('recommendation', []), fill=LIGHT_BLUE)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('input_json')
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    data = json.loads(Path(args.input_json).read_text(encoding='utf-8'))
    build(data, args.output)

if __name__ == '__main__':
    main()
