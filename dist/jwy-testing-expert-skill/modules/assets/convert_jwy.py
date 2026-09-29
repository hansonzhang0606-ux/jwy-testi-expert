# -*- coding: utf-8 -*-
"""
JWY 专用 XMind(经典 content.xml) -> DMP Excel 转换器。

针对 Step3-选项1(xmind-testcase) 产出的 5 层脑图结构：
  L1 = 模块(一/二/三/四/五...)
  L2 = 子模块
  L3 = 条件/场景 (C3 第1参)   -> 作为用例名称来源
  L4 = 数据   (C3 第2参)      -> input 步骤
  L5 = 预期   (C3 第3参)      -> output 预期
「需求疑问点(待确认)」分支为说明性叶子，不生成用例。

遵循 xmind-to-testcase SKILL 约束：
  - 基于团队 DMP 模板(excel模板.xlsx)，前 4 行样式原样保留
  - 第 3 行英文列名建立 col_map，按名映射，绝不硬编码列号
  - 第 5 行起写数据
  - caseGroup(功能路径) = 版本衍生"年-年月-年月日" + 故事名称(全表同一值)
  - manager(责任人) = 当前会话测试人员
  - version(适用版本) = 测试人员告知的 V...
  - team/product/modulePath = 固定值({部署配置}，沿用模板示例)
  - name 自动加"验证"前缀
"""
import argparse, sys, zipfile, re
from openpyxl import load_workbook
from openpyxl.styles import Alignment

NS = "{urn:xmind:xmap:xmlns:content:2.0}"

def title(e):
    t = e.find(f"{NS}title")
    return t.text if t is not None else ""

def kids(e):
    ch = e.find(f"{NS}children")
    if ch is None:
        return []
    ts = ch.find(f"{NS}topics")
    if ts is None:
        return []
    return ts.findall(f"{NS}topic")

def walk(e, depth, out, skip_branch):
    """收集 L3(条件) -> L4(数据) -> L5(预期) 链。"""
    if skip_branch:
        return
    tg = title(e)
    if depth == 1 and tg == "需求疑问点(待确认)":
        return  # 说明性分支，整体跳过
    ks = kids(e)
    if depth == 3:
        # L3 条件节点：期望一个 L4(数据) 含一个 L5(预期)
        data = title(ks[0]) if len(ks) >= 1 else ""
        expect = title(kids(ks[0])[0]) if (len(ks) >= 1 and kids(ks[0])) else ""
        out.append({"condition": tg, "data": data, "expect": expect})
        return
    for c in ks:
        walk(c, depth + 1, out, skip_branch=False)

def parse_xmind(path):
    with zipfile.ZipFile(path) as z:
        content = z.read("content.xml").decode("utf-8")
    import xml.etree.ElementTree as ET
    root = ET.fromstring(content)
    sheet = root.find(f"{NS}sheet")
    root_topic = sheet.find(f"{NS}topic")
    out = []
    for l1 in kids(root_topic):
        walk(l1, 1, out, skip_branch=False)
    return out

def split_items(text):
    if not text:
        return []
    text = str(text).replace("；", ";").replace("；", ";").replace("\n", ";")
    items = [t.strip(" 。.；;") for t in text.split(";")]
    return [t for t in items if t]

def build_io(data, expect):
    data_items = split_items(data)
    expect_items = split_items(expect)
    inp, out = [], []
    inp.append("进入对应功能页面，准备测试账号与测试数据")
    out.append("页面正常加载，无报错，测试数据就绪")
    n = max(len(data_items), len(expect_items))
    for i in range(n):
        t = data_items[i] if i < len(data_items) else "复核上述操作"
        e = expect_items[i] if i < len(expect_items) else "系统返回符合预期"
        inp.append("执行操作：%s" % t)
        out.append("预期结果：%s" % e)
    inp.append("检查系统响应、页面展示与数据落库结果")
    out.append("系统与页面展示与预期一致，数据正确落库")
    while len(inp) < 3:
        inp.append("执行测试用例步骤并观察结果")
        out.append("结果符合预期")
    return "\n".join("%d. %s" % (i + 1, s) for i, s in enumerate(inp)), \
           "\n".join("%d. %s" % (i + 1, s) for i, s in enumerate(out))

def case_level(condition):
    if any(k in condition for k in ["异常", "错误", "失败", "无效", "为空", "缺失"]):
        return "P1"
    if any(k in condition for k in ["边界", "临界", "最大", "最小", "重复", "超长"]):
        return "P2"
    return "P0"

def version_to_path(version, story_name):
    """V2026.8.25 -> 2026-202608-20260825-{story_name}"""
    m = re.match(r"^V?(\d{4})\.?(\d{1,2})\.?(\d{1,2})", version.strip())
    if not m:
        # 兜底：直接用版本原文
        return f"{version.strip()}-{story_name}"
    y, mo, d = m.group(1), m.group(2).zfill(2), m.group(3).zfill(2)
    return f"{y}-{y}{mo}-{y}{mo}{d}-{story_name}"

def convert(xmind, template, output, version, manager, story_name,
           team, product, module_path):
    recs = parse_xmind(xmind)
    wb = load_workbook(template)
    ws = wb.active
    col_map = {}
    for c in range(1, ws.max_column + 1):
        k = ws.cell(row=3, column=c).value
        if k and str(k).strip():
            col_map[str(k).strip()] = c
    # 清空第 5 行起
    for r in range(5, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None

    case_group = version_to_path(version, story_name)
    data_start = 5
    for i, r in enumerate(recs):
        row = data_start + i
        inp, out = build_io(r["data"], r["expect"])
        tc = {
            "team": team,
            "caseGroup": case_group,
            "name": "验证" + r["condition"],
            "preCondition": "1. 测试环境已部署\n2. 中鼎征信接口/RPA 通道相关依赖服务正常运行\n3. 测试账号与测试数据就绪",
            "input": inp,
            "output": out,
            "product": product,
            "modulePath": module_path,
            "version": version,
            "caseType": "功能测试",
            "source": "",
            "caseLevel": case_level(r["condition"]),
            "manager": manager,
            "autoState": "否",
        }
        for col_name, val in tc.items():
            if col_name in col_map:
                ws.cell(row=row, column=col_map[col_name], value=val)
        for cn in ("preCondition", "input", "output"):
            if cn in col_map:
                ws.cell(row=row, column=col_map[cn]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 120
    wb.save(output)
    return len(recs), case_group

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--xmind", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--version", required=True)
    ap.add_argument("--manager", required=True)
    ap.add_argument("--story-name", required=True)
    ap.add_argument("--team", default="{部署配置}")
    ap.add_argument("--product", default="{部署配置}")
    ap.add_argument("--module-path", default="{部署配置}")
    args = ap.parse_args()
    n, cg = convert(args.xmind, args.template, args.output, args.version,
                    args.manager, args.story_name, args.team, args.product, args.module_path)
    print(f"已生成 {n} 条用例 -> {args.output}")
    print(f"caseGroup(功能路径) = {cg}")
    print(f"manager(责任人) = {args.manager}  version(适用版本) = {args.version}")
